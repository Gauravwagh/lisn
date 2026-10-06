"""Regenerate the binary fixtures: python tests/fixtures/make_fixtures.py"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent

LEFT_COLUMN = [
    "The first column starts here with a sen-",
    "tence that is hyphenated across lines.",
    "It continues in the left column and",
    "ends properly.",
]
RIGHT_COLUMN = [
    "The right column comes second in the",
    "reading order. It has its own sentence.",
]


def make_pdf() -> None:
    import pymupdf

    doc = pymupdf.open()
    for page_number in (1, 2):
        page = doc.new_page(width=400, height=500)
        page.insert_text((40, 25), "Running Header Of The Report", fontsize=8)  # header zone
        if page_number == 1:
            page.insert_text((40, 70), "Big Title", fontsize=20)
        y = 110
        for line in LEFT_COLUMN:
            page.insert_text((30, y), line, fontsize=9)
            y += 12
        y = 110
        for line in RIGHT_COLUMN:
            page.insert_text((210, y), line, fontsize=9)
            y += 12
        page.insert_text(
            (30, 200), f"Full width paragraph on page {page_number} that spans the whole page width.", fontsize=9
        )
        page.insert_text((30, 215), "Footnote marker", fontsize=9)
        page.insert_text((108, 211), "1", fontsize=5)  # superscript-ish, same line
        page.insert_text((190, 485), str(page_number), fontsize=8)  # page number in footer zone
    doc.set_metadata({"title": "Fixture Report"})
    doc.save(str(HERE / "sample.pdf"))


def make_docx() -> None:
    import docx

    document = docx.Document()
    document.add_heading("Fixture Document", level=0)
    document.add_paragraph("An opening paragraph with two sentences. Here is the second one.")
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("Chapter one body text.")
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("Second bullet", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "1"
    document.add_heading("Chapter Two", level=1)
    document.add_paragraph("Chapter two body text.")
    document.save(str(HERE / "sample.docx"))


def make_epub() -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("lisn-fixture")
    book.set_title("Fixture Book")
    book.set_language("en")
    chapters = []
    for number, (title, body) in enumerate(
        [("One", "The first chapter has a sentence. And another one."), ("Two", "The second chapter is brief.")], 1
    ):
        chapter = epub.EpubHtml(title=f"Chapter {title}", file_name=f"chap_{number}.xhtml", lang="en")
        chapter.content = (
            f"<html><body><h1>Chapter {title}</h1><p>{body}</p><p>Footnote<sup>1</sup> here.</p></body></html>"
        )
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *chapters]
    epub.write_epub(str(HERE / "sample.epub"), book)


def make_html() -> None:
    html = """<!doctype html><html><head><title>Fixture Article</title></head><body>
<nav><a href="/">Home</a> <a href="/about">About</a></nav>
<article>
<h1>Fixture Article</h1>
<p>The article body has a first paragraph with enough words to be considered real content by the extractor.</p>
<h2>A Section</h2>
<p>The section paragraph also has a reasonable number of words so that it is kept as main content.</p>
<ul><li>First item in a list</li><li>Second item in a list</li></ul>
</article>
<aside>Related links and advertising noise that should be removed.</aside>
<footer>Copyright 2024 Fixture Inc. All rights reserved.</footer>
</body></html>"""
    (HERE / "sample.html").write_text(html, encoding="utf-8")


if __name__ == "__main__":
    make_pdf()
    make_docx()
    make_epub()
    make_html()
    print("fixtures written to", HERE)
